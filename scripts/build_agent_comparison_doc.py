#!/usr/bin/env python3
"""生成《与材料学智能体的定位对比与深度评估》DOCX + Markdown 两个版本。

文章来源是单一内容结构（CONTENT），两个渲染器共用，避免文档与 md 漂移。

用法：
    python scripts/build_agent_comparison_doc.py
产物：
    docs/submission/2026-10-29-与材料学智能体对比-定位差异与深度评估.docx
    docs/submission/2026-10-29-与材料学智能体对比-定位差异与深度评估.md
"""

from __future__ import annotations

import os
import sys

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT_DIR = os.path.join(ROOT, "docs", "submission")
BASENAME = "2026-10-29-与材料学智能体对比-定位差异与深度评估"
DOCX_PATH = os.path.join(OUT_DIR, BASENAME + ".docx")
MD_PATH = os.path.join(OUT_DIR, BASENAME + ".md")

FONT_CN = "Microsoft YaHei"
FONT_CODE = "Consolas"
NAVY = RGBColor(0x1F, 0x2A, 0x44)
BLUE = RGBColor(0x1F, 0x38, 0x64)
GREY = RGBColor(0x59, 0x59, 0x59)
USABLE_CM = 16.4


# --------------------------------------------------------------------------
# 内容（唯一事实来源）
# --------------------------------------------------------------------------

TITLE = "与材料学智能体对比：定位、重合度与深度"
SUBTITLE = "ChemCrtic ｜ 可比性判定系统 ｜ 2026-10-29 ｜ 赛前内部参考"

CONTENT = [
    ("h1", "结论速览（30 秒）"),
    ("bullets", [
        "对方是执行型智能体：跑 DFT / MD、查材料数据库、批量工作流，产出的是一手数据。",
        "我方是判定型系统：判定两条数值能否同框，产出证书与补证排序，本身不产数据。",
        "唯一的重合在「跨来源数值比对」，但对象不同——对方解释差异（计算方法），我方拒绝并列（测量口径）。",
        "深度：执行深度我方明显不如；判定深度我方更专，但只覆盖一个反应、一种量。",
    ]),

    ("h1", "0 · 本文口径与限制"),
    ("p", "本文用于赛前对外表述与答辩准备。对比对象是 SciencePal 平台内置的默认智能体「材料学智能体」"
          "（sciencepal.ai/agents/detail/material，界面标注 v1.1.1）；对比方是本次参赛作品 ChemCrtic。"),
    ("note", "限制：该页面是前端渲染的应用，服务端抓取只返回页面标题。本文的能力明细来自两处——"
             "① 页面「01 可以开展哪些工作」的六项能力（截图逐条核对）；"
             "② 其余五节只能读到标题（02 从问题到交付物 / 03 成果可以继续使用 / 04 从一个具体任务开始 / "
             "05 支撑工作的工具 / 06 这个智能体的演进）。凡涉及工具链的表述均标注为「推断」。"),
    ("p", "我方数据基线：参照集 32 条已核验文献记录（全部带 locator）；判定引擎 chemcritic-engine/1.1.0，"
          "自检 26 项全过。"),

    ("h1", "1 · 对方是什么：一双能真跑计算的手"),
    ("lead", "它从结构出发，把物理计算跑完，再从真实输出文件里取数。"),
    ("table", {
        "widths": [3.7, 7.2, 5.5],
        "headers": ["能力（原文）", "实际在做什么", "可信度来自哪里"],
        "rows": [
            ["晶体结构与缺陷", "CIF/POSCAR/XYZ 转换、对称性分析、表面与缺陷构建，准备计算结构", "结构库（推断：pymatgen / ASE 一类）"],
            ["电子结构与原子模拟", "设置 FHI-aims DFT 或 LAMMPS 模拟、检查收敛、从真实输出文件提取性质", "物理引擎 + 输出文件解析"],
            ["跨数据库材料证据", "通过 OPTIMADE 与 Materials Project 跨库比对同一材料记录，保留查询信息并说明数值差异", "OPTIMADE 客户端 + MP API"],
            ["批量工作流与筛选", "用 atomate2 / FireWorks 组织弛豫、弹性、声子、状态方程、缺陷任务，隔离失败任务", "工作流引擎（jobflow）"],
            ["实验数据分析", "解析 XRD 与电化学数据，拟合阻抗或其他曲线", "谱图与拟合脚本（推断）"],
            ["材料机器学习", "构建描述符与性质模型；使用机器学习原子间势", "训练与推理栈（推断）"],
        ],
    }),
    ("p", "它的「真」不来自模型：收敛就是收敛，数值是从真实输出文件里抽出来的。"
          "这一点与我方「判定不来自模型」是同一种诚实，只是路径不同——它靠物理引擎，我们靠规则引擎。"),

    ("h1", "2 · 我方是什么：一张会拒答的嘴"),
    ("lead", "它判定两个数字能不能放进同一句话；不能，就给出阻断维度与解锁条件。"),
    ("table", {
        "widths": [2.6, 8.2, 5.6],
        "headers": ["层", "做什么", "落在平台的哪里"],
        "rows": [
            ["记录层", "文献 → entry 级结构化记录：口径、校准、原文措辞、locator、DOI", "个人知识库 + 沙箱"],
            ["判定层", "8 维门槛（硬 4 / 软 4）在沙箱内确定性执行，逐对输出证书", "沙箱代码执行 + Skill"],
            ["行动层", "补证杠杆（补哪一条解锁最多）、档内定位、跨档迁移目标", "引擎输出（本次新增）"],
        ],
    }),
    ("p", "模型只能把证书翻译成人话，不能改写 verdict。这一条是产品红线，不是实现细节。"),

    ("h1", "3 · 逐维对比"),
    ("table", {
        "widths": [3.2, 6.6, 6.6],
        "headers": ["维度", "材料学智能体", "ChemCrtic"],
        "rows": [
            ["领域", "无机材料 / 计算材料", "有机合成（Suzuki–Miyaura 偶联）"],
            ["输入", "结构文件、计算输出、数据库记录", "已核验的文献数值 + 一次实验记录"],
            ["输出", "结构、收敛的计算、工作流、拟合曲线、模型", "可比性证书、阻断对、补证排序"],
            ["真值来源", "物理引擎 + 真实输出文件", "确定性规则引擎 + 已核验记录"],
            ["判定形态", "给结论，并解释数值差异", "拒绝并列，给阻断维度与解锁条件"],
            ["证据链", "查询 provenance", "逐维 locator / DOI / 原文措辞 / 校准级别"],
            ["可复现", "重跑即复现", "重跑引擎即复现（记录依赖人工核验）"],
            ["成熟度", "v1.1.1，已上线产品", "6 天原型（无训练组件、32 条记录）"],
        ],
    }),

    ("h1", "4 · 重合度：表面约三成，判定层不重合"),
    ("p", "会让人产生「这不就是一个东西」的只有两处：它的「03 跨数据库材料证据」对应我们的多源数值比对；"
          "它的「05 实验数据分析」对应我们的取证与实验设计。但两边的对象不同——这是全部差异的根。"),
    ("table", {
        "widths": [3.0, 6.2, 7.2],
        "headers": ["对比项", "材料学智能体 03", "ChemCrtic"],
        "rows": [
            ["比的量", "材料属性的计算值", "化学产率的实验测量值"],
            ["差异来源", "方法：泛函、赝势、收敛、超胞", "口径：分离 / GC / NMR / 转化率、内标、校准"],
            ["能否消除", "能——同一参数重算即可", "不能——重做实验也消不掉口径差"],
            ["因此的动作", "解释差异，然后继续给结论", "拒绝并列，给阻断维度与解锁条件"],
        ],
    }),
    ("lead", "它解释差异，我们拒绝并列；它产出数据，我们审计数据。"),
    ("h2", "同一个问题的两种回答"),
    ("p", "问：这个底物的收率大概是多少？"),
    ("table", {
        "widths": [3.4, 13.0],
        "headers": ["回答者", "给出的答案"],
        "rows": [
            ["通用文献助手", "35% – 100%"],
            ["材料学智能体", "会给出数据库/计算值的横向对比，并解释数值差异（推断）"],
            ["ChemCrtic", "8 个报告值里只有 3 个能进同一等价类；其余分别卡在口径（GLC / NMR / 分离）、校准与体系"],
        ],
    }),
    ("note", "这 8 个数值全部来自已核验文献（SM-001…SM-006），不是构造的例子："
             "66%（GLC）与同篇 40%（分离）判 BLOCKED，卡在维度 3（口径不同，禁止换算）；"
             "只写「35%」而不说怎么测的，判 NEEDS_EVIDENCE，卡在维度 4（口径标记缺原话支撑）。"),

    ("h1", "5 · 深度：两种深度，不在同一根轴上"),
    ("table", {
        "widths": [4.3, 3.4, 3.3, 5.4],
        "headers": ["轴", "材料学智能体", "ChemCrtic", "依据"],
        "rows": [
            ["执行深度（能否真跑计算）", "9 / 10", "2 / 10", "我方只有纯 Python 规则引擎与物料衡算，不跑 DFT/MD、不解析谱图"],
            ["判定深度（能否同框）", "3 / 10", "8 / 10", "对方 03 只到「说明差异」；我方输出 verdict / 阻断维度 / 解锁条件"],
            ["证据链颗粒度", "7 / 10", "8 / 10", "对方保留查询 provenance；我方逐维 locator + 校准 + 原文措辞"],
            ["覆盖广度", "9 / 10", "2 / 10", "我方限定一个反应、一种量"],
            ["可复现", "9 / 10", "7 / 10", "我方记录依赖人工核验"],
            ["成熟度", "9 / 10", "4 / 10", "对方已上线；我方无训练组件"],
        ],
    }),
    ("p", "结论要说清楚：执行深度我们明显不如；判定深度我们更专，但更窄。"
          "这不是「同一产品谁做得更好」，是「手」与「闸」的区别。"),
    ("p", "评分口径：10 分制，按「该维度能否在与本产品同类的任务上独立完成并给出可核对结果」打分，"
          "不是按工程量打分。"),

    ("h1", "6 · 我方新增的差异化：从「拒答」到「排序」"),
    ("p", "对方 03 只说「两个库的数值为什么不同」，不给「先补哪一条」。这正是我方最近一次迭代补上的能力："
          "把「不能比」换算成「补哪一条最值」。"),
    ("code",
     "27 对不能并列：3 对可以靠补证据解决，24 对是确定不同（补也没用）\n"
     "补 1 条 → 解锁 3 对\n"
     "1. 回原文确认该条记录的口径与校准（含内标）（新解锁 3 对，累计 3）\n"
     "\n"
     "—— 语料级（补哪几条记录的口径，能新增最多可比对）\n"
     "496 个两两比较：38 对可直接并列；36 对只缺口径确认\n"
     "只补第一条记录的口径（SM-004.1）→ 新增 13 对可比\n"
     "补这 3 条 → 可比对 38 → 74"),
    ("table", {
        "widths": [5.4, 4.6, 6.4],
        "headers": ["指标", "数值", "口径"],
        "rows": [
            ["参照集两两比较", "496 对", "32 条记录的全组合"],
            ["可直接并列", "38 对（7.7%）", "通过全部硬门槛"],
            ["仅缺口径确认", "36 对", "硬门槛 UNKNOWN，且无 FAIL"],
            ["补 3 条记录后", "74 对", "上限值：假设补上的证据能通过门槛"],
        ],
    }),
    ("note", "限定语（对外必须带）：以上数字的适用范围是「我们的 32 条参照集 + 当前规则版本」，"
             "不是化学界的统计结论。补证解锁数是上限估计，不得说成承诺。"),

    ("h1", "7 · 对外三问三答"),
    ("h2", "问：这不就是跨数据库比较吗？"),
    ("p", "答：它比的是计算值，差异来自方法，重算能消；我们比的是实验值，差异来自口径，重算消不掉，只能靠门槛隔离。"),
    ("h2", "问：他也能解释数值差异。"),
    ("p", "答：解释不等于禁止。我们的 verdict 由代码给出，带阻断维度与解锁条件；它不会拦住任何一次并列，我们专门用于拦截。"),
    ("h2", "问：深度不如他。"),
    ("p", "答：承认执行深度不如。他在全域更广，我们在单点更深——这是两种产品，不是同一产品的两代。"),

    ("h1", "8 · 风险与边界（如实披露）"),
    ("bullets", [
        "我方无学习组件（没有模型训练）。FPR / FNR / BDA / CAL 只证明规则内部一致，不证明泛化能力。",
        "参照集 32 条；同底物 + 同口径 + 可校准的最多 5 条，统计上撑不住「领域结论」，只能作为方法可行性证据。",
        "对 PDF 知识库的自主检索尚未端到端验收，当前流程是「结构化记录 + 沙箱引擎」。",
        "对方是已上线产品（v1.1.1），工具链整合度、覆盖面均高于我方；不应在任何材料中暗示我方具备同类执行能力。",
    ]),

    ("h1", "9 · 结论与三条建议"),
    ("bullets", [
        "定位写成「文献数值的可比性判定」，不写「文献检索」——前者是判定，后者是检索，是两个产品。",
        "对外叙事先讲现象（八个报告值拼不出一个区间），不要用竞品开场；用竞品定义自己会天然显得衍生。",
        "差异化只认一条：拒绝并列 + 补证排序。执行深度不争，争了就输。",
    ]),

    ("h1", "附录 A · 采集与方法"),
    ("bullets", [
        "页面抓取：https://sciencepal.ai/agents/detail/material → HTTP 200，正文仅返回标题（前端渲染）。",
        "能力明细：页面「01 可以开展哪些工作」六项，逐条核对截图；其余五节仅读到标题。",
        "我方数据：data/records.json（32 条）、单文件引擎 --leverage 输出、--selftest（26 passed, 0 failed）。",
        "本文生成日期：2026-10-29。若对方页面内容更新，第 1、3、4 节的判断需重新核对。",
    ]),
]


# --------------------------------------------------------------------------
# DOCX 渲染
# --------------------------------------------------------------------------

def _set_cn_font(run, name: str = FONT_CN) -> None:
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:eastAsia"), name)
    rfonts.set(qn("w:ascii"), name)
    rfonts.set(qn("w:hAnsi"), name)


def _shade(cell, hex_fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tc_pr.append(shd)


def _para(doc, text: str, size: float, bold: bool = False, color=None,
          space_before: float = 0, space_after: float = 6, italic: bool = False,
          line: float = 1.35, keep_next: bool = False):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(space_before)
    pf.space_after = Pt(space_after)
    pf.line_spacing = line
    if keep_next:
        pf.keep_with_next = True
    run = p.add_run(text)
    run.bold = bold
    run.italic = italic
    run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = color
    _set_cn_font(run)
    return p


def _add_table(doc, spec: dict) -> None:
    headers = spec["headers"]
    rows = spec["rows"]
    widths = spec["widths"]

    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    tbl = table._tbl
    tbl_pr = tbl.tblPr
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tbl_pr.append(layout)
    tbl_w = OxmlElement("w:tblW")
    tbl_w.set(qn("w:w"), str(sum(int(round(w * 567)) for w in widths)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_pr.append(tbl_w)
    for grid_col, width in zip(tbl.tblGrid.findall(qn("w:gridCol")), widths):
        grid_col.set(qn("w:w"), str(int(round(width * 567))))

    hdr = table.rows[0].cells
    for i, text in enumerate(headers):
        hdr[i].width = Cm(widths[i])
        _shade(hdr[i], "E9EEF7")
        p = hdr[i].paragraphs[0]
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        run = p.add_run(text)
        run.bold = True
        run.font.size = Pt(9)
        run.font.color.rgb = BLUE
        _set_cn_font(run)

    for row in rows:
        cells = table.add_row().cells
        for i, text in enumerate(row):
            cells[i].width = Cm(widths[i])
            p = cells[i].paragraphs[0]
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.15
            run = p.add_run(text)
            run.font.size = Pt(9)
            _set_cn_font(run)

    # 表后留白
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(4)
    spacer.paragraph_format.line_spacing = 1.0
    for run in spacer.runs:
        run.font.size = Pt(4)


def _add_code(doc, text: str) -> None:
    for line in text.split("\n"):
        p = doc.add_paragraph()
        pf = p.paragraph_format
        pf.space_before = Pt(0)
        pf.space_after = Pt(0)
        pf.left_indent = Cm(0.6)
        pf.line_spacing = 1.15
        p_pr = p._p.get_or_add_pPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), "F4F4F6")
        p_pr.append(shd)
        run = p.add_run(line if line else " ")
        run.font.name = FONT_CODE
        run.font.size = Pt(8.5)
        run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), FONT_CN)
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(6)
    spacer.paragraph_format.line_spacing = 1.0


def _add_footer(sec, text: str) -> None:
    p = sec.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    run.font.size = Pt(8)
    run.font.color.rgb = GREY
    _set_cn_font(run)

    field_run = p.add_run()
    field_run.font.size = Pt(8)
    field_run.font.color.rgb = GREY
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    sep = OxmlElement("w:fldChar")
    sep.set(qn("w:fldCharType"), "separate")
    placeholder = OxmlElement("w:t")
    placeholder.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for node in (begin, instr, sep, placeholder, end):
        field_run._r.append(node)

    tail = p.add_run(" 页")
    tail.font.size = Pt(8)
    tail.font.color.rgb = GREY
    _set_cn_font(tail)


def build_docx() -> str:
    doc = Document()

    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(2.3)
    sec.top_margin = Cm(2.2)
    sec.bottom_margin = Cm(2.0)

    normal = doc.styles["Normal"]
    normal.font.name = FONT_CN
    normal.font.size = Pt(10.5)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), FONT_CN)

    _add_footer(sec, "ChemCrtic · 与材料学智能体对比 · 第 ")

    _para(doc, TITLE, 19, bold=True, color=NAVY, space_after=4, line=1.2)
    _para(doc, SUBTITLE, 9.5, color=GREY, space_after=14)

    for block in CONTENT:
        kind = block[0]
        if kind == "h1":
            _para(doc, block[1], 14, bold=True, color=BLUE, space_before=14, space_after=6,
                  line=1.2, keep_next=True)
        elif kind == "h2":
            _para(doc, block[1], 11, bold=True, color=NAVY, space_before=8, space_after=3,
                  line=1.25, keep_next=True)
        elif kind == "p":
            _para(doc, block[1], 10.5)
        elif kind == "lead":
            _para(doc, block[1], 11, bold=True, color=NAVY, space_before=2, space_after=8,
                  line=1.3, keep_next=True)
        elif kind == "note":
            _para(doc, block[1], 9.5, color=GREY, italic=True, space_after=8, line=1.3)
        elif kind == "bullets":
            for item in block[1]:
                p = doc.add_paragraph(style="List Bullet")
                p.paragraph_format.space_after = Pt(3)
                p.paragraph_format.line_spacing = 1.3
                run = p.add_run(item)
                run.font.size = Pt(10.5)
                _set_cn_font(run)
        elif kind == "table":
            _add_table(doc, block[1])
        elif kind == "code":
            _add_code(doc, block[1])
        else:  # pragma: no cover
            raise ValueError(f"unknown block: {kind}")

    doc.core_properties.title = TITLE
    doc.core_properties.subject = "ChemCrtic 与 SciencePal 内置材料学智能体的对比"
    doc.core_properties.author = "ChemCrtic"
    doc.core_properties.comments = "赛前内部参考；页面采集限制见附录 A"

    os.makedirs(OUT_DIR, exist_ok=True)
    doc.save(DOCX_PATH)
    return DOCX_PATH


# --------------------------------------------------------------------------
# Markdown 渲染
# --------------------------------------------------------------------------

def build_md() -> str:
    out = [f"# {TITLE}", "", f"*{SUBTITLE}*", ""]
    for block in CONTENT:
        kind = block[0]
        if kind == "h1":
            out += [f"## {block[1]}", ""]
        elif kind == "h2":
            out += [f"### {block[1]}", ""]
        elif kind == "p":
            out += [block[1], ""]
        elif kind == "lead":
            out += [f"**{block[1]}**", ""]
        elif kind == "note":
            out += [f"> {block[1]}", ""]
        elif kind == "bullets":
            out += [f"- {item}" for item in block[1]] + [""]
        elif kind == "table":
            spec = block[1]
            out.append("| " + " | ".join(spec["headers"]) + " |")
            out.append("|" + "|".join(["---"] * len(spec["headers"])) + "|")
            for row in spec["rows"]:
                out.append("| " + " | ".join(row) + " |")
            out.append("")
        elif kind == "code":
            out += ["```text", block[1], "```", ""]
    os.makedirs(OUT_DIR, exist_ok=True)
    with open(MD_PATH, "w", encoding="utf-8", newline="\n") as fh:
        fh.write("\n".join(out).rstrip() + "\n")
    return MD_PATH


def main() -> int:
    docx_path = build_docx()
    md_path = build_md()
    for path in (docx_path, md_path):
        print(f"generated {os.path.relpath(path, ROOT)}  ({os.path.getsize(path)} bytes)")
    return 0


if __name__ == "__main__":
    sys.exit(main())
