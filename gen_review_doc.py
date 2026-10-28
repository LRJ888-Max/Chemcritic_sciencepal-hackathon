"""生成 ChemCritic 项目三维评判报告 Word 文档。"""
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn

OUT = r"C:\Users\Lenovo\Desktop\sciencepal\ChemCritic项目评判报告.docx"

doc = Document()

# 页面设置
for section in doc.sections:
    section.top_margin = Cm(2.5)
    section.bottom_margin = Cm(2.5)
    section.left_margin = Cm(2.8)
    section.right_margin = Cm(2.8)

# 默认字体
style = doc.styles['Normal']
style.font.name = '微软雅黑'
style.font.size = Pt(11)
style.element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑')


def set_cn_font(run, name='微软雅黑', size=11, bold=False, color=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn('w:eastAsia'), name)
    run.font.size = Pt(size)
    run.font.bold = bold
    if color:
        run.font.color.rgb = color


def add_title(text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text)
    set_cn_font(r, size=18, bold=True)
    p.paragraph_format.space_after = Pt(6)


def add_subtitle(text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text)
    set_cn_font(r, size=10, color=RGBColor(0x66, 0x66, 0x66))
    p.paragraph_format.space_after = Pt(18)


def add_h1(text):
    p = doc.add_paragraph()
    r = p.add_run(text)
    set_cn_font(r, size=14, bold=True, color=RGBColor(0x1F, 0x3A, 0x5F))
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(8)


def add_h2(text):
    p = doc.add_paragraph()
    r = p.add_run(text)
    set_cn_font(r, size=12, bold=True)
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(4)


def add_body(text, bold=False):
    p = doc.add_paragraph()
    r = p.add_run(text)
    set_cn_font(r, size=11, bold=bold)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.4
    return p


def add_bullet(text):
    p = doc.add_paragraph(style='List Bullet')
    r = p.add_run(text)
    set_cn_font(r, size=11)
    p.paragraph_format.space_after = Pt(2)
    return p


def add_table(headers, rows, col_widths=None):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = 'Light Grid Accent 1'
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = ''
        p = cell.paragraphs[0]
        r = p.add_run(h)
        set_cn_font(r, size=10, bold=True)
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            cell = table.rows[ri + 1].cells[ci]
            cell.text = ''
            p = cell.paragraphs[0]
            r = p.add_run(str(val))
            set_cn_font(r, size=10)
    if col_widths:
        for i, w in enumerate(col_widths):
            for row in table.rows:
                row.cells[i].width = Cm(w)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


# ===== 正文开始 =====

add_title('ChemCritic 项目三维评判报告')
add_subtitle('科学严谨性 · 创新性 · AI技术性 ｜ 2026-10-28')

# 概述
add_h1('一、项目概述')
add_body('ChemCritic 是 SciencePal AI4Science Hackathon 2026 赛道 4（AI Agent）的参赛作品，定位为'
         '“实验没做出来之后，告诉你下一个该做哪个实验的多智能体系统”。核心场景为 Suzuki 偶联反应，'
         '通过 8 维门槛规则引擎判断不同文献中的收率数字能否同框比较，再给出档内定位与下一步实验建议。')
add_body('本评判基于对核心源码（engine/gates.py、placement.py、schemas.py）、评测设计（eval/goldset.json、'
         'baseline.md）及数据规模的实际阅读，而非仅依据 README。')

# ===== 科学严谨性 =====
add_h1('二、科学严谨性：7 / 10')

add_h2('做对的地方')
add_table(
    ['维度', '证据'],
    [
        ['问题真实且锋利', '抓住“假精确”现象——8 个报告值里只有 3 个构成同一等价类，有文献 DOI 支撑'],
        ['判定可审计', '硬门槛走确定性 Python，LLM 无权改写判定结果'],
        ['UNKNOWN 纪律严格', '原文未报告字段一律记 UNKNOWN，不补猜；用户没给数值直接返回 NEEDS_EVIDENCE'],
        ['证据分级', 'evidence_grade 分 A_confirmed / candidate / imported_unchecked；label_raw 要求原文逐字支撑'],
        ['同篇不独立', '_same_source() 防止同一篇文献的不同 entry 被当作两条独立来源'],
    ],
    col_widths=[4, 12]
)

add_h2('硬伤')
add_bullet('底物/产物比对目前是字符串匹配——_canon() 只做 lowercase+去空格。'
           '“4-bromoanisole”和“1-bromo-4-methoxybenzene”会被判为不同底物。代码 TODO 自己也写了要接 rdkit，'
           '但现在没接。这在化学上是致命的。')
add_bullet('样本量太小——已核验记录仅 6 条（SM-001~006），34 条候选池未核验。'
           'Suzuki 反应 40 年文献量是天文数字，6 条记录算覆盖率没有统计意义。')
add_bullet('阈值拍脑袋——温度差 ≥20°C 标 L2，代码注释写“A 于 10/25 定：Suzuki 里 20°C 已足以显著改变转化率”，'
           '但这个阈值的化学文献依据未给出。')
add_bullet('档位划分主观——A/B/C 三档本质是关键词匹配（profiles.py 仅 2KB），没有严格的化学分类学依据。')
add_bullet('Gold set 仅 20 对，其中 12 对（G01~G12）还是从引擎自身派生的 seed 对——'
           '文件 warning 里自己承认“FPR/FNR 会偏乐观”，真正独立标注的只有 8 对。')

# ===== 创新性 =====
add_h1('三、创新性：6.5 / 10')

add_h2('亮点')
add_bullet('切入点选得聪明——不做又一个“文献检索 Agent”或“反应预测模型”，'
           '而是瞄准“数字能不能同框比较”这个被忽视的问题。Demo 故事好讲：“你的实验没失败，这个体系天花板就是 40%”。')
add_bullet('“档内定义异常”的 framing 有价值——不是泛泛说“35% 偏低”，而是“在你的体系档位里正常区间是多少”，'
           '对实验化学家有实际参考意义。')

add_h2('不足')
add_bullet('核心常识门槛低——“isolated yield 和 GC yield 不能直接比”是有经验的化学家本来就知道的事。'
           '目标用户到底是谁？新手化学家需要的可能不是这个。')
add_bullet('技术上零突破——没有新模型、新算法、新架构。本质是精心设计的 prompt workflow + 确定性规则脚本。'
           '黑客松里够用，说是“创新”就勉强了。')
add_bullet('最有想象力的部分还没做——A4“序贯判别实验决策树”是 README 吹得最响的，但目前是 🚧。'
           '如果 A3/A4 做不出来，整个产品退化成一个“文献口径检查器”。')

# ===== AI技术性 =====
add_h1('四、AI技术性：6 / 10')

add_h2('工程素养是好的')
add_bullet('单文件 bundle 生成（build_bundle.py）、零依赖引擎、19 项自检脚本、数据结构清晰')
add_bullet('对 LLM 能力边界有清醒认识——明确写“LLM 只能解释、不能改写判定”')
add_bullet('防御性编程做得扎实——到处查 None、空列表、边界情况')

add_h2('但“AI 含量”低')
add_table(
    ['对外说法', '实际情况'],
    [
        ['六个多 Agent 系统', '核心价值全在 A2 一个确定性模块里。A3（机制推演）、A4（实验设计）还没实现，真正“智能”的部分恰恰是空缺的'],
        ['Subagent 并行推演', '就是 fan-out 多个 LLM 调用，没有技术创新'],
        ['知识库', '11 篇 PDF，规模极小'],
        ['闭环记忆', 'A5 还是 🚧'],
    ],
    col_widths=[4, 12]
)
add_body('一句话：这是一个工程完成度高、但 AI 技术深度浅的项目。它的核心竞争力不在 AI，在化学规则设计。', bold=True)

# ===== 总评 =====
add_h1('五、总评与行动建议')

add_table(
    ['维度', '分数', '一句话'],
    [
        ['科学严谨性', '7 / 10', '意识很好，但深度、样本量、化学结构识别都不够'],
        ['创新性', '6.5 / 10', '切入点新，但技术无突破，最有价值的部分还没做'],
        ['AI 技术性', '6 / 10', '工程扎实，但“多智能体”是包装，核心是规则引擎'],
    ],
    col_widths=[4, 2.5, 9.5]
)

add_h2('语境判断')
add_bullet('黑客松语境下：选题聪明、demo 故事好讲、比大部分套壳作品强，能打。')
add_bullet('放到严肃研究里：样本量、化学深度、A3/A4 落地程度都撑不起一个完整产品的叙事。')

add_h2('明天演示前最该补的一件事')
add_body('把 rdkit canonical SMILES 接上，哪怕是简化版。现在底物比对靠字符串匹配，'
         '评审里如果有懂化学的人，这是第一个会被问的问题。', bold=True)

doc.save(OUT)
print(f"Saved: {OUT}")
